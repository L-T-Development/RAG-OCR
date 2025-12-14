import fitz  # PyMuPDF
import difflib
from docx import Document
from openpyxl import load_workbook


# ---------- COMMON HELPERS ----------

def normalize_lines(text: str):
    """
    Simple normalization.
    Keeps logic intentionally weak.
    """
    lines = text.splitlines()
    return [line.strip().lower() for line in lines if line.strip()]


def diff_lines(old_lines, new_lines):
    diff = difflib.ndiff(old_lines, new_lines)

    added = []
    removed = []
    modified = []

    for line in diff:
        if line.startswith("+ "):
            added.append(line[2:])
        elif line.startswith("- "):
            removed.append(line[2:])
        elif line.startswith("? "):
            modified.append(line[2:])

    return {
        "added": added,
        "removed": removed,
        "modified": modified
    }


# ---------- PDF ----------

def extract_pdf(path):
    doc = fitz.open(path)
    text = ""
    for page in doc:
        text += page.get_text()
    return normalize_lines(text)


def compare_pdfs(old_pdf, new_pdf):
    return diff_lines(
        extract_pdf(old_pdf),
        extract_pdf(new_pdf)
    )


# ---------- DOCX ----------

def extract_docx(path):
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    return normalize_lines(text)


def compare_docx(old_docx, new_docx):
    return diff_lines(
        extract_docx(old_docx),
        extract_docx(new_docx)
    )


# ---------- EXCEL ----------

def compare_excels(old_xlsx, new_xlsx):
    wb_old = load_workbook(old_xlsx)
    wb_new = load_workbook(new_xlsx)

    changes = []

    for sheet in wb_old.sheetnames:
        if sheet not in wb_new.sheetnames:
            continue

        ws_old = wb_old[sheet]
        ws_new = wb_new[sheet]

        for row in ws_old.iter_rows():
            for cell in row:
                old_val = cell.value
                new_val = ws_new[cell.coordinate].value
                if old_val != new_val:
                    changes.append(
                        f"{sheet} {cell.coordinate}: {old_val} → {new_val}"
                    )

    return {
        "added": [],
        "removed": [],
        "modified": changes
    }
