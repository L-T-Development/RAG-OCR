import fitz  # PyMuPDF
import difflib
from docx import Document
from openpyxl import load_workbook
import os
import uuid
import time
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from typing import Dict, Any, Optional


# ---------------- COMMON HELPERS ----------------

def extract_lines(text: str):
    """
    Extract lines preserving original content (not normalized).
    Returns list of lines with whitespace stripped.
    """
    lines = text.splitlines()
    return [line.strip() for line in lines if line.strip()]


def unified_line_diff(old_lines, new_lines):
    """
    Generate unified line-by-line diff with status for each line.
    Returns dict with lines, stats, and total_lines
    Status: 'equal', 'added', 'removed', 'modified'
    """
    diff_result = []
    
    # Use difflib to get unified diff
    matcher = difflib.SequenceMatcher(None, old_lines, new_lines)
    
    line_num = 1
    stats = {"added": 0, "removed": 0, "modified": 0, "equal": 0}
    
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == 'equal':
            # Lines are the same
            for i in range(i1, i2):
                diff_result.append({
                    "line": line_num,
                    "status": "equal",
                    "text": old_lines[i],
                    "old_text": None,
                    "new_text": None
                })
                stats["equal"] += 1
                line_num += 1
                
        elif tag == 'replace':
            # Lines are different (modified)
            max_lines = max(i2 - i1, j2 - j1)
            for k in range(max_lines):
                old_text = old_lines[i1 + k] if (i1 + k) < i2 else None
                new_text = new_lines[j1 + k] if (j1 + k) < j2 else None
                
                if old_text and new_text:
                    diff_result.append({
                        "line": line_num,
                        "status": "modified",
                        "text": new_text,
                        "old_text": old_text,
                        "new_text": new_text
                    })
                    stats["modified"] += 1
                elif old_text:
                    diff_result.append({
                        "line": line_num,
                        "status": "removed",
                        "text": old_text,
                        "old_text": old_text,
                        "new_text": None
                    })
                    stats["removed"] += 1
                elif new_text:
                    diff_result.append({
                        "line": line_num,
                        "status": "added",
                        "text": new_text,
                        "old_text": None,
                        "new_text": new_text
                    })
                    stats["added"] += 1
                    
                line_num += 1
                
        elif tag == 'delete':
            # Lines removed
            for i in range(i1, i2):
                diff_result.append({
                    "line": line_num,
                    "status": "removed",
                    "text": old_lines[i],
                    "old_text": old_lines[i],
                    "new_text": None
                })
                stats["removed"] += 1
                line_num += 1
                
        elif tag == 'insert':
            # Lines added
            for j in range(j1, j2):
                diff_result.append({
                    "line": line_num,
                    "status": "added",
                    "text": new_lines[j],
                    "old_text": None,
                    "new_text": new_lines[j]
                })
                stats["added"] += 1
                line_num += 1

    return {
        "lines": diff_result,
        "stats": stats,
        "total_lines": len(diff_result)
    }


# ---------------- PDF ----------------

def extract_pdf(path):
    print(f"[PDF] Extracting text from: {path}")
    doc = fitz.open(path)
    text_lines = []
    
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text()
        
        # Split text into lines and clean them
        raw_lines = text.splitlines()
        for line in raw_lines:
            cleaned_line = line.strip()
            if cleaned_line:  # Only keep non-empty lines
                text_lines.append(cleaned_line)
    
    print(f"[PDF] Total lines extracted: {len(text_lines)}")
    doc.close()
    return text_lines





def compare_pdfs(old_pdf, new_pdf):
    """Compare PDFs line by line"""
    old_lines = extract_pdf(old_pdf)
    new_lines = extract_pdf(new_pdf)
    return unified_line_diff(old_lines, new_lines)











# ---------------- DOCX ----------------

def extract_docx(path):
    print(f"[DOCX] Extracting text from: {path}")
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    lines = extract_lines(text)
    print(f"[DOCX] Extracted {len(lines)} lines")
    return lines


def compare_docx(old_docx, new_docx):
    old_lines = extract_docx(old_docx)
    new_lines = extract_docx(new_docx)
    return unified_line_diff(old_lines, new_lines)


# ---------------- EXCEL ----------------

def extract_excel(path):
    """Extract Excel content as lines (row by row)"""
    print(f"[XLSX] Extracting content from: {path}")
    wb = load_workbook(path)
    lines = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"[Sheet: {sheet_name}]")
        
        for row in ws.iter_rows(values_only=True):
            # Convert row to string, filtering None values
            row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if row_str.strip() and row_str.replace("|", "").replace(" ", ""):
                lines.append(row_str)
    
    print(f"[XLSX] Extracted {len(lines)} lines")
    return lines


def compare_excels(old_xlsx, new_xlsx):
    old_lines = extract_excel(old_xlsx)
    new_lines = extract_excel(new_xlsx)
    return unified_line_diff(old_lines, new_lines)


# ============================================================================
# BACKGROUND JOB MANAGEMENT (ThreadPoolExecutor for non-blocking comparison)
# ============================================================================

# ThreadPoolExecutor for background comparison jobs (max 3 concurrent workers)
_COMPARISON_EXECUTOR = ThreadPoolExecutor(max_workers=3, thread_name_prefix="CompareWorker")

# In-memory job storage with thread-safe access
_jobs_lock = threading.Lock()
_jobs_store: Dict[str, dict] = {}


def create_comparison_job(old_path: str, new_path: str, old_filename: str, new_filename: str, file_ext: str) -> str:
    """
    Create a new background comparison job.
    
    Args:
        old_path: Path to old file (temporary)
        new_path: Path to new file (temporary)
        old_filename: Original filename of old file
        new_filename: Original filename of new file
        file_ext: File extension (.pdf, .docx, .xlsx)
    
    Returns:
        job_id: Unique identifier for tracking this job
    
    Thread-safe: Uses lock when modifying jobs store
    """
    job_id = str(uuid.uuid4())
    
    # Create job record
    job = {
        "job_id": job_id,
        "status": "pending",  # pending | processing | completed | failed
        "created_at": datetime.now().isoformat(),
        "started_at": None,
        "finished_at": None,
        "progress": 0,  # 0-100
        "old_filename": old_filename,
        "new_filename": new_filename,
        "file_extension": file_ext,
        "old_path": old_path,
        "new_path": new_path,
        "result": None,
        "error": None
    }
    
    # Store job in thread-safe manner
    with _jobs_lock:
        _jobs_store[job_id] = job
    
    # Submit to executor (non-blocking)
    _COMPARISON_EXECUTOR.submit(_run_comparison_worker, job_id)
    
    print(f"[JOB] Created job {job_id} for comparing {old_filename} vs {new_filename}")
    return job_id


def get_comparison_status(job_id: str) -> Optional[Dict[str, Any]]:
    """
    Retrieve current status of a comparison job.
    
    Args:
        job_id: Job identifier
    
    Returns:
        Job status dict (without internal file paths) or None if not found
    
    Thread-safe: Uses lock for reading
    """
    with _jobs_lock:
        job = _jobs_store.get(job_id)
        if not job:
            return None
        
        # Return copy without internal paths
        public_job = {k: v for k, v in job.items() if k not in ['old_path', 'new_path']}
        return public_job


def _run_comparison_worker(job_id: str):
    """
    Worker function that performs the actual comparison in background thread.
    
    This runs asynchronously - the API returns immediately while this processes.
    All existing comparison logic remains UNCHANGED.
    
    Args:
        job_id: Job to process
    
    Thread-safe: Updates job state with lock protection
    """
    print(f"[WORKER] Starting comparison job {job_id}")
    start_time = time.perf_counter()
    
    # Get job details (thread-safe read)
    with _jobs_lock:
        job = _jobs_store.get(job_id)
        if not job:
            print(f"[WORKER] ERROR: Job {job_id} not found!")
            return
    
    try:
        # Update status to processing
        with _jobs_lock:
            job["status"] = "processing"
            job["started_at"] = datetime.now().isoformat()
            job["progress"] = 10
        
        print(f"[WORKER] Processing: {job['old_filename']} vs {job['new_filename']}")
        print(f"[WORKER] File type: {job['file_extension']}")
        
        # ============================================================
        # STEP 1: Run comparison (EXISTING LOGIC - NO CHANGES)
        # ============================================================
        with _jobs_lock:
            job["progress"] = 30
        
        file_ext = job["file_extension"]
        old_path = job["old_path"]
        new_path = job["new_path"]
        
        if file_ext == ".pdf":
            diff_result = compare_pdfs(old_path, new_path)
        elif file_ext == ".docx":
            diff_result = compare_docx(old_path, new_path)
        elif file_ext == ".xlsx":
            diff_result = compare_excels(old_path, new_path)
        else:
            raise ValueError(f"Unsupported file type: {file_ext}")
        
        stats = diff_result.get('stats', {})
        print(f"[WORKER] Diff complete - Added: {stats.get('added', 0)}, "
              f"Removed: {stats.get('removed', 0)}, "
              f"Modified: {stats.get('modified', 0)}, "
              f"Unchanged: {stats.get('equal', 0)}")
        
        # ============================================================
        # STEP 2: Generate LLM summary (EXISTING LOGIC - NO CHANGES)
        # ============================================================
        with _jobs_lock:
            job["progress"] = 70
        
        print(f"[WORKER] Generating LLM summary...")
        # Import here to avoid circular dependency
        from process.llm_summary import summarize_diff
        summary = summarize_diff(diff_result)
        if not summary:
            print(f"[WORKER] WARNING: LLM summary failed, using fallback")
            summary = "LLM unavailable. Raw diff available in response."
        
        # ============================================================
        # STEP 3: Store results
        # ============================================================
        processing_time = round(time.perf_counter() - start_time, 3)
        
        with _jobs_lock:
            job["status"] = "completed"
            job["finished_at"] = datetime.now().isoformat()
            job["progress"] = 100
            job["result"] = {
                "summary": summary,
                "diff": diff_result,
                "processing_time_seconds": processing_time
            }
        
        print(f"[WORKER] OK Job {job_id} completed in {processing_time}s")
    
    except Exception as e:
        # Handle errors gracefully (thread-safe)
        print(f"[WORKER] ERROR in job {job_id}: {type(e).__name__}: {str(e)}")
        import traceback
        traceback.print_exc()
        
        with _jobs_lock:
            job["status"] = "failed"
            job["finished_at"] = datetime.now().isoformat()
            job["error"] = str(e)
    
    finally:
        # Cleanup temporary files
        for path in [job.get("old_path"), job.get("new_path")]:
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                    print(f"[CLEANUP] Removed temp file: {path}")
                except Exception as e:
                    print(f"[CLEANUP] WARNING: Failed to remove {path}: {e}")
