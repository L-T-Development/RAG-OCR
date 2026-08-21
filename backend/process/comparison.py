"""
Cross-document column comparison — the ONE entry point every caller uses.
================================================================================

Before this module existed there were two chat comparison paths with different
matching rules: `views.chat_thread` (Router 1) used `table_search_engine` — schema
resolution, ConfirmedMatch, boilerplate filtering, deterministic likely-matches —
while `pipeline/query.py` (Router 2) used an older fuzzy engine in
`pipeline/tables.py` (SequenceMatcher ≥ 0.75), which merges sequential part numbers.
Router 1 also silently compared only the first two of N named files.

Everything now funnels through here:

    resolve_documents()      names → Document rows in the thread's scope (any extension)
    compare_documents()      two documents, one concept   → engine result dict
    compare_many()           one source vs N targets      → presence matrix
    format_*()               markdown for chat
    attach_excel_report()    Excel + download link (report store lives in views)

The engine itself is `TableSearchEngine.compare_tables` (table_search_engine.py);
this module only decides *what* to compare and renders the answer.
"""
from __future__ import annotations

import os
import re
from typing import Any, Dict, Iterable, List, Optional

from django.db.models import Q

from .models import Document
from .table_search_engine import compare_structured_sources, compare_pdfs

# Documents that can carry extracted tables (see pipeline/ingestion._DISPATCH).
TABLE_EXTENSIONS = (".pdf", ".xlsx", ".xls", ".docx")

_FILE_MENTION_RE = re.compile(r"@?([\w\-. ]+\.(?:pdf|xlsx|xls|docx))", re.IGNORECASE)


# ── Scope & document resolution ────────────────────────────────────────────────

def thread_scope_ids(thread) -> List[str]:
    """This thread plus every ancestor (sub-threads inherit their parents' documents)."""
    ids, cur = [str(thread.id)], thread
    while cur.parent:
        ids.append(str(cur.parent.id))
        cur = cur.parent
    return ids


def documents_in_scope(thread):
    """Table-bearing documents visible from `thread` (own + inherited), upload order."""
    q = Q()
    for ext in TABLE_EXTENSIONS:
        q |= Q(filename__iendswith=ext)
    return (Document.objects
            .filter(thread_id__in=thread_scope_ids(thread))
            .filter(q)
            .order_by("uploaded_at"))


def mentioned_files(query: str) -> List[str]:
    """Filenames named in a query (`@a.pdf`, or bare `a.pdf`), deduped, in order."""
    seen, out = set(), []
    for m in _FILE_MENTION_RE.findall(query):
        name = m.strip()
        key = name.lower()
        if key not in seen:
            seen.add(key)
            out.append(name)
    return out


def resolve_documents(thread, names: Iterable[str]) -> List[Document]:
    """Map filenames (as typed) to Document rows: exact (case-insensitive) first,
    then substring. Unresolvable names are skipped; order and uniqueness kept."""
    docs = list(documents_in_scope(thread))
    out: List[Document] = []
    for name in names:
        n = name.strip().lower()
        hit = next((d for d in docs if d.filename.lower() == n), None)
        if not hit:
            hit = next((d for d in docs if n in d.filename.lower()), None)
        if hit and hit not in out:
            out.append(hit)
    return out


# ── The comparison itself ──────────────────────────────────────────────────────

def compare_documents(doc_a: Document, doc_b: Document, column: str,
                      column_b: Optional[str] = None,
                      comparison_type: str = "all",
                      match_any_column_in_b: bool = False,
                      prefer_literal: bool = False,
                      thread=None) -> Dict[str, Any]:
    """
    Compare one concept (or two explicitly named columns) between two documents.

    1. DB-backed: persisted ExtractedTable rows (works after the upload is gone,
       and for xlsx/docx which have no PDF to fall back to).
    2. File-backed fallback for PDFs still on disk, forcing a fresh extraction so
       the tables the DB path just cached under the same filename are not reused.

    Returns the engine result dict; `found` is False with `error` (+ available
    columns) when the column could not be resolved, or `error == "no_data"` when
    neither path had tables. `engine_path` records which path answered.
    """
    thread_ids = thread_scope_ids(thread) if thread is not None else None
    kwargs = dict(
        column_name_pdf2=column_b,
        match_any_column_in_pdf2=match_any_column_in_b,
        comparison_type=comparison_type,
        category1=doc_a.category,
        category2=doc_b.category,
        prefer_literal=prefer_literal,
    )

    result = compare_structured_sources(
        doc_a.filename, doc_b.filename, column,
        doc1_id=str(doc_a.id), doc2_id=str(doc_b.id), thread_ids=thread_ids,
        **kwargs,
    )
    if result.get("found"):
        result["engine_path"] = "db"
        return result
    db_result = result

    a_pdf = doc_a.filename.lower().endswith(".pdf") and _file_exists(doc_a)
    b_pdf = doc_b.filename.lower().endswith(".pdf") and _file_exists(doc_b)
    if a_pdf and b_pdf:
        result = compare_pdfs(doc_a.file.path, doc_b.file.path, column,
                              force_extract=True,
                              doc1_id=str(doc_a.id), doc2_id=str(doc_b.id), **kwargs)
        result["engine_path"] = "file"
        if result.get("found"):
            return result
        # Prefer whichever failure carries the real header lists (helps the
        # "column not found" message); the file pass usually sees more of them.
        if not result.get("available_columns_pdf1") and db_result.get("available_columns_pdf1"):
            result = db_result
        return result

    # No file to fall back to (xlsx/docx, or the PDF was not retained on disk).
    if not db_result.get("error") or "not available" in str(db_result.get("error")):
        db_result["error"] = "no_data"
        missing = [d.filename for d, ok in ((doc_a, a_pdf), (doc_b, b_pdf))
                   if d.filename.lower().endswith(".pdf") and not ok]
        db_result["missing_files"] = missing
    db_result["engine_path"] = "db"
    return db_result


def _file_exists(doc: Document) -> bool:
    try:
        return bool(doc.file) and os.path.exists(doc.file.path)
    except Exception:
        return False


def compare_many(source: Document, targets: List[Document], column: str,
                 prefer_literal: bool = False, thread=None) -> Dict[str, Any]:
    """
    One source vs N targets ("are the ISPL's part numbers in each of these
    files?"). Runs the pairwise engine per target so every rule (schema,
    ConfirmedMatch, boilerplate, likely-matches) applies identically, then folds
    the results into a per-value presence matrix from the source's point of view.

    matrix[value][target] ∈ {"exact", "confirmed", "likely", "missing", "n/a"}
    ("n/a" = that target could not be compared: no tables / column not found).
    """
    per_target: Dict[str, Dict[str, Any]] = {}
    errors: Dict[str, str] = {}
    source_values: set = set()
    matrix: Dict[str, Dict[str, str]] = {}
    likely_notes: Dict[str, Dict[str, str]] = {}   # value → target → "≈ other value (reason)"

    for t in targets:
        r = compare_documents(source, t, column, comparison_type="all",
                              prefer_literal=prefer_literal, thread=thread)
        per_target[t.filename] = r
        if not r.get("found"):
            errors[t.filename] = _short_error(r)
            continue
        res = r["results"]
        common = set(res.get("common", {}).get("values", []))
        missing = set(res.get("in_pdf1_only", {}).get("values", []))
        excluded = set(res.get("excluded", {}).get("values", []))
        likely = {p["pdf1_value"]: p for p in res.get("likely_matches", {}).get("pairs", [])}
        confirmed = {p["pdf1_value"]: p for p in res.get("confirmed_matches", {}).get("pairs", [])}
        vals = common | missing | set(likely) | set(confirmed)
        source_values |= vals
        for v in vals:
            row = matrix.setdefault(v, {})
            if v in common:
                row[t.filename] = "exact"
            elif v in confirmed:
                row[t.filename] = "confirmed"
                likely_notes.setdefault(v, {})[t.filename] = f"≡ {confirmed[v]['pdf2_value']}"
            elif v in likely:
                row[t.filename] = "likely"
                likely_notes.setdefault(v, {})[t.filename] = (
                    f"≈ {likely[v]['pdf2_value']} ({likely[v]['reason']})")
            else:
                row[t.filename] = "missing"
        # values excluded as boilerplate are dropped from the matrix entirely
        for v in excluded:
            matrix.pop(v, None)
            source_values.discard(v)

    ok_targets = [t.filename for t in targets if t.filename not in errors]
    for v, row in matrix.items():
        for name in errors:
            row.setdefault(name, "n/a")

    def _accounted(status):
        return status in ("exact", "confirmed", "likely")

    missing_everywhere = sorted(v for v, row in matrix.items()
                                if ok_targets and all(row.get(n) == "missing" for n in ok_targets))
    present_everywhere = sorted(v for v, row in matrix.items()
                                if ok_targets and all(_accounted(row.get(n)) for n in ok_targets))
    partial = sorted(v for v, row in matrix.items()
                     if v not in missing_everywhere and v not in present_everywhere)

    # Column label: what the source resolved to (same for every target).
    resolved = next((r.get("column", "") for r in per_target.values() if r.get("found")), column)
    src_label = resolved.split(" ↔ ")[0].strip() if " ↔ " in resolved else resolved

    per_target_counts = {}
    for name in ok_targets:
        col = [row.get(name) for row in matrix.values()]
        per_target_counts[name] = {
            "exact": col.count("exact"), "confirmed": col.count("confirmed"),
            "likely": col.count("likely"), "missing": col.count("missing"),
            "resolved_column": (per_target[name].get("column", "").split(" ↔ ")[-1].strip()
                                if " ↔ " in per_target[name].get("column", "") else per_target[name].get("column", "")),
        }

    return {
        "found": bool(ok_targets),
        "source": source.filename,
        "targets": [t.filename for t in targets],
        "ok_targets": ok_targets,
        "errors": errors,
        "column": src_label,
        "source_total": len(matrix),
        "matrix": matrix,
        "notes": likely_notes,
        "missing_everywhere": missing_everywhere,
        "present_everywhere": present_everywhere,
        "partial": partial,
        "per_target": per_target_counts,
        "per_target_results": per_target,
    }


def _short_error(r: Dict[str, Any]) -> str:
    err = r.get("error") or "comparison failed"
    if err == "no_data":
        return "no extracted tables (re-upload, or run reextract_tables)"
    return str(err)


# ── Text mode: a column vs another document's whole text ───────────────────────

_TEXT_MODE_PHRASES = (
    "mentioned", "mention", "anywhere", "appear in", "appears in", "referenced",
    "referred", "in the text", "full text", "anywhere in", "covered in", "listed in",
)


def wants_text_mode(query: str) -> bool:
    """True when the user asked about *mentions* rather than a column-to-column diff
    ("are all these parts mentioned anywhere in the manual?")."""
    q = (query or "").lower()
    return any(p in q for p in _TEXT_MODE_PHRASES)


def collect_column_values(doc: Document, column: str, thread=None,
                          prefer_literal: bool = False) -> Dict[str, Any]:
    """
    Distinct values of one concept/column in one document, from its persisted
    tables. Returns {found, values, rows, resolved_column, available_columns}.

    This is the source side of a text-mode comparison: we still need a real column
    in the spares list, we just no longer need one in the target.
    """
    import pandas as pd
    from .column_roles import resolve as resolve_column
    from .field_schema import normalize_value, is_empty_value, is_boilerplate_value
    from .table_search_engine import _search_engine

    thread_ids = thread_scope_ids(thread) if thread is not None else None
    dfs = _search_engine.load_structured_tables(
        doc.filename, doc_id=str(doc.id), thread_ids=thread_ids)
    if not dfs:
        return {"found": False, "error": "no_data", "values": set(), "rows": [],
                "resolved_column": None, "available_columns": []}

    available, values, rows, resolved = set(), set(), [], None
    excluded = set()
    for df in dfs:
        data_cols = [c for c in df.columns if not str(c).startswith("_")]
        available.update(str(c) for c in data_cols)
        col = resolve_column(column, data_cols, doc.category, doc_id=str(doc.id),
                             prefer_literal=prefer_literal)
        if not col:
            continue
        resolved = resolved or col
        for _, row in df.iterrows():
            raw = str(row[col]) if pd.notna(row[col]) else ""
            if is_empty_value(raw):
                continue
            nv = normalize_value(raw)
            if is_boilerplate_value(nv):
                excluded.add(nv)
                continue
            values.add(nv)
            rec = {c: (str(row[c]) if pd.notna(row[c]) else "")
                   for c in df.columns if not str(c).startswith("_")}
            rec["_matched_value"] = raw.strip()
            rec["_norm_value"] = nv
            if "_source_page" in row:
                rec["_page"] = int(row["_source_page"])
            rows.append(rec)

    return {
        "found": bool(values),
        "error": None if values else "column_not_found",
        "values": values, "rows": rows, "excluded": excluded,
        "resolved_column": resolved,
        "available_columns": sorted(available),
    }


def compare_against_text(source: Document, targets: List[Document], column: str,
                         thread=None, prefer_literal: bool = False) -> Dict[str, Any]:
    """
    "Is every value of this column mentioned ANYWHERE in these documents?"

    Unlike the column-to-column path, the targets need no comparable column and no
    tables at all — each source value is looked up in the target's indexed text
    (prose + table cells, see process/mentions.py) and reported with the pages it
    was found on. This is the only path that can answer a spares-list-vs-manual
    question, because manuals mention parts in running text.
    """
    from . import mentions

    src = collect_column_values(source, column, thread=thread, prefer_literal=prefer_literal)
    if not src["found"]:
        return {"found": False, "error": src["error"], "source": source.filename,
                "available_columns": src["available_columns"]}

    values = sorted(src["values"])
    indexed = mentions.indexed_documents([str(t.id) for t in targets])
    per_target, errors = {}, {}
    hits_by_target: Dict[str, Dict[str, List[dict]]] = {}
    for t in targets:
        if mentions.doc_key(t.id) not in indexed:
            errors[t.filename] = "not indexed yet — run `manage.py build_mention_index`"
            continue
        hits_by_target[t.filename] = mentions.lookup_values(values, str(t.id))

    ok_targets = [t.filename for t in targets if t.filename in hits_by_target]
    matrix: Dict[str, Dict[str, Any]] = {}
    for v in values:
        row = {}
        for name in ok_targets:
            hits = hits_by_target[name].get(v) or []
            row[name] = {
                "found": bool(hits),
                "near": False, "variant": None, "reason": None,
                "pages": sorted({h["page"] for h in hits})[:10],
                "kinds": sorted({h["kind"] for h in hits}),
            }
        matrix[v] = row

    # Near-match tier — the same deterministic rules the column comparison uses
    # (leading zeros, annotation prefix/suffix, typos in word-like values). Without
    # this, text mode reports "XL17461 NAMICA" as absent from a manual that says
    # "XL17461", which contradicts the column answer for the very same documents.
    from .field_schema import pair_likely_values
    for t in targets:
        name = t.filename
        if name not in hits_by_target:
            continue
        unmatched = [v for v in values if not matrix[v][name]["found"]]
        if not unmatched:
            continue
        vocab = mentions.vocabulary(str(t.id))
        if not vocab:
            continue
        for a, b, reason in pair_likely_values(set(unmatched), set(vocab)):
            hits = mentions.pages_for(b, str(t.id))
            cell = matrix[a][name]
            cell.update({
                "found": True, "near": True, "variant": b, "reason": reason,
                "pages": sorted({h["page"] for h in hits})[:10],
                "kinds": sorted({h["kind"] for h in hits}),
            })

    missing_everywhere = [v for v in values
                          if ok_targets and not any(matrix[v][n]["found"] for n in ok_targets)]
    found_everywhere = [v for v in values
                        if ok_targets and all(matrix[v][n]["found"] for n in ok_targets)]
    partial = [v for v in values if v not in missing_everywhere and v not in found_everywhere]
    near_anywhere = [v for v in values
                     if any(matrix[v][n].get("near") for n in ok_targets)]

    for name in ok_targets:
        found = [v for v in values if matrix[v][name]["found"]]
        in_text = [v for v in found if "text" in matrix[v][name]["kinds"]]
        near = [v for v in found if matrix[v][name].get("near")]
        per_target[name] = {
            "found": len(found), "missing": len(values) - len(found),
            "via_prose": len(in_text), "near": len(near),
        }

    row_by_value = {}
    for r in src["rows"]:
        row_by_value.setdefault(r["_norm_value"], r)

    return {
        "found": bool(ok_targets),
        "mode": "text",
        "source": source.filename,
        "targets": [t.filename for t in targets],
        "ok_targets": ok_targets,
        "errors": errors,
        "column": src["resolved_column"] or column,
        "source_total": len(values),
        "values": values,
        "matrix": matrix,
        "rows": row_by_value,
        "missing_everywhere": missing_everywhere,
        "found_everywhere": found_everywhere,
        "partial": partial,
        "near_anywhere": near_anywhere,
        "per_target": per_target,
        "excluded": sorted(src.get("excluded") or []),
    }


def format_text_mode_response(m: Dict[str, Any]) -> str:
    """Chat rendering of compare_against_text()."""
    src, col = m["source"], m["column"]
    ok = m["ok_targets"]
    n = len(ok)
    out = (f"## Mention check: {src} vs {len(m['targets'])} document"
           f"{'s' if len(m['targets']) != 1 else ''}\n\n")
    out += (f"**Column checked:** `{col}` in {src}\n\n"
            f"_Each value is searched in the **full text** of the target — paragraphs, "
            f"notes and table cells alike — not just a matching column. Use this when the "
            f"target is a manual rather than a spares table._\n\n")

    if m["errors"]:
        out += "> ⚠️ Could not search: " + "; ".join(
            f"**{k}** ({v})" for k, v in m["errors"].items()) + "\n\n"
    if not ok:
        return out + "None of the target documents could be searched."

    out += "| | Count |\n|---|---:|\n"
    out += f"| Distinct `{col}` values in {src} | {m['source_total']} |\n"
    out += f"| ✅ Mentioned in **all** {n} document{'s' if n != 1 else ''} | {len(m['found_everywhere'])} |\n"
    if m.get("near_anywhere"):
        out += (f"| ⚠️ — of those, matched only as a variant (review) "
                f"| {len(m['near_anywhere'])} |\n")
    if m["partial"]:
        out += f"| ◐ Mentioned in some, missing from others | {len(m['partial'])} |\n"
    out += f"| ❌ Not mentioned **anywhere** | {len(m['missing_everywhere'])} |\n\n"

    out += ("**Per document:**\n\n"
            "| Document | Mentioned | Of which in prose | Variant match | Not found |\n"
            "|---|---:|---:|---:|---:|\n")
    for name in ok:
        c = m["per_target"][name]
        out += (f"| {name} | {c['found']} | {c['via_prose']} "
                f"| {c.get('near', 0)} | {c['missing']} |\n")
    out += "\n---\n\n"

    if m["missing_everywhere"]:
        out += f"### ❌ Not mentioned anywhere ({len(m['missing_everywhere'])})\n\n"
        for i, v in enumerate(m["missing_everywhere"][:25], 1):
            r = m["rows"].get(v, {})
            desc = next((str(val) for k, val in r.items()
                         if not k.startswith("_") and val
                         and any(w in k.lower() for w in
                                 ("nomenclature", "description", "designation", "name"))), "")
            out += f"{i}. `{v}`" + (f" — {desc[:60]}" if desc else "") + "\n"
        if len(m["missing_everywhere"]) > 25:
            out += f"\n_...and {len(m['missing_everywhere']) - 25} more — full list in the Excel report._\n"
        out += "\n---\n\n"

    shown = (m["partial"] or m["found_everywhere"])[:20]
    if shown:
        title = ("◐ Mentioned in some documents only" if m["partial"]
                 else "✅ Where each value was found")
        out += f"### {title}\n\n| Value | " + " | ".join(ok) + " |\n|---|" + "---|" * n + "\n"
        for v in shown:
            cells = []
            for name in ok:
                cell = m["matrix"][v][name]
                if not cell["found"]:
                    cells.append("❌")
                else:
                    pages = ", ".join(str(p) for p in cell["pages"][:4])
                    more = "…" if len(cell["pages"]) > 4 else ""
                    if cell.get("near"):
                        cells.append(f"⚠️ as `{cell['variant']}` p{pages}{more}")
                    else:
                        cells.append(f"✅ p{pages}{more}")
            out += f"| `{v}` | " + " | ".join(cells) + " |\n"
        total = len(m["partial"] or m["found_everywhere"])
        if total > len(shown):
            out += f"\n_...and {total - len(shown)} more — full list in the Excel report._\n"
    if m.get("excluded"):
        out += (f"\n_{len(m['excluded'])} generic placeholder value(s) "
                f"(e.g. 'Standard Item') were not searched._\n")
    return out


def generate_text_mode_excel(m: Dict[str, Any]) -> bytes:
    """Per-value × per-document mention report, with the pages each hit was on."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    bad_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    def _header(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(horizontal="center")

    ws = wb.active
    ws.title = "Summary"
    _header(ws, ["Metric", "Value"])
    ws.append(["Mode", "Mention check (source column vs target full text)"])
    ws.append(["Source file", m["source"]])
    ws.append(["Column", m["column"]])
    ws.append(["Target documents", ", ".join(m["targets"])])
    ws.append(["Distinct values searched", m["source_total"]])
    ws.append(["Mentioned in all documents", len(m["found_everywhere"])])
    ws.append(["Mentioned in some only", len(m["partial"])])
    ws.append(["Not mentioned anywhere", len(m["missing_everywhere"])])
    for name, err in m["errors"].items():
        ws.append([f"Not searched: {name}", err])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80

    ok = m["ok_targets"]
    ws = wb.create_sheet("Mentions")
    _header(ws, [m["column"]] + [f"{n} (pages)" for n in ok] + ["Status"])
    order = m["missing_everywhere"] + m["partial"] + m["found_everywhere"]
    for v in order:
        cells = []
        for name in ok:
            cell = m["matrix"][v][name]
            if not cell["found"]:
                cells.append("not found")
            else:
                pg = ", ".join(f"p{p}" for p in cell["pages"])
                cells.append(f"as {cell['variant']} ({pg})" if cell.get("near") else pg)
        status = ("not mentioned anywhere" if v in m["missing_everywhere"]
                  else "mentioned everywhere" if v in m["found_everywhere"] else "partial")
        ws.append([v] + cells + [status])
        r = ws.max_row
        for j, name in enumerate(ok, start=2):
            ws.cell(row=r, column=j).fill = ok_fill if m["matrix"][v][name]["found"] else bad_fill
    ws.column_dimensions["A"].width = 28
    for j in range(2, len(ok) + 3):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 32
    ws.freeze_panes = "B2"

    ws = wb.create_sheet("Not Mentioned")
    _header(ws, [m["column"], "Row context from " + m["source"]])
    for v in m["missing_everywhere"]:
        r = m["rows"].get(v, {})
        ctx = "; ".join(f"{k}: {val}" for k, val in r.items()
                        if not k.startswith("_") and str(val).strip())[:500]
        ws.append([v, ctx])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Rendering ──────────────────────────────────────────────────────────────────

def format_available_columns(columns, limit: int = 20) -> str:
    """Render the header list for the "column not found" message.

    The list arrives sorted alphabetically, which buries real headers behind
    "Column_7" placeholders and stray numeric values from tables whose header row
    was never recovered — so a truncated list showed the user nothing usable.
    Real labels come first, and the count of what was dropped is stated.
    """
    cols = [str(c) for c in columns if str(c).strip()]
    if not cols:
        return ''

    def _is_placeholder(c):
        return bool(re.fullmatch(r'Column_\d+(_\d+)?', c)) or not any(ch.isalpha() for ch in c)

    named = [c for c in cols if not _is_placeholder(c)]
    placeholders = [c for c in cols if _is_placeholder(c)]
    shown = (named + placeholders)[:limit]
    text = ', '.join(shown)
    hidden = len(cols) - len(shown)
    if hidden > 0:
        text += f" … (+{hidden} more)"
    if placeholders and not named:
        text += ("\n\n> ⚠️ No column names were recovered from this file — its table "
                 "headers did not extract. Re-upload or reprocess it.")
    return text


def suggest_identifier_columns(doc: Document, thread=None) -> List[tuple]:
    """Columns in `doc` whose *values* look like identifiers — a hint for the
    "Column Not Found" message when the header is one we have never seen."""
    try:
        from .column_profile import columns_from_frames, identifier_columns
        from .table_search_engine import _search_engine
        dfs = _search_engine.load_structured_tables(
            doc.filename, doc_id=str(doc.id),
            thread_ids=thread_scope_ids(thread) if thread is not None else None)
        if not dfs:
            return []
        return identifier_columns(columns_from_frames(dfs))
    except Exception as e:
        print(f"[COMPARISON] column suggestion failed: {e}")
        return []


def format_column_not_found(result: Dict[str, Any], requested: str,
                            name_a: str, name_b: str,
                            suggestions: Optional[Dict[str, List[tuple]]] = None) -> str:
    """The "we found tables but not that column" answer, listing real headers so
    the user can name one exactly (quoted) or fix the schema override."""
    cols1 = format_available_columns(result.get('available_columns_pdf1', []))
    cols2 = format_available_columns(result.get('available_columns_pdf2', []))
    msg = (
        "## ❌ Column Not Found for Comparison\n\n"
        f"Requested column: **{requested}**\n\n"
        f"{result.get('error', '')}\n\n"
        f"**{name_a} columns:** {cols1 or 'N/A'}\n\n"
        f"**{name_b} columns:** {cols2 or 'N/A'}\n\n"
    )
    hints = []
    for fname in (name_a, name_b):
        for header, why in (suggestions or {}).get(fname, [])[:3]:
            hints.append(f"- `{header}` in **{fname}** — {why}")
    if hints:
        msg += ("**These columns hold identifier-shaped values**, judged from the data "
                "rather than the header — one of them may be the column you meant:\n"
                + "\n".join(hints) + "\n\n")
    msg += (
        "Name the column explicitly in quotes, e.g. "
        f"`compare \"<header in {name_a}>\" in @{name_a} with \"<header in {name_b}>\" in @{name_b}`.\n\n"
        "If this document type always labels the column that way, add the header to "
        "`field_schema.json` (see `GET /api/field-schema/`) so it resolves automatically."
    )
    return msg


def format_no_data(result: Dict[str, Any], name_a: str, name_b: str) -> str:
    missing = result.get("missing_files") or []
    if missing:
        return (
            "## ❌ Assessment Could Not Be Completed\n\n"
            "Structured tables are not available and source PDF file(s) are missing:\n"
            f"- {', '.join(missing)}\n\n"
            "Please re-upload and reprocess these files, then run:\n"
            "`Compare DRG No present or not between @file1.pdf and @file2.pdf`"
        )
    return (
        "## ❌ No Extracted Tables\n\n"
        f"Neither **{name_a}** nor **{name_b}** has extracted tables to compare. "
        "Make sure both finished processing; documents uploaded before table "
        "extraction was added can be back-filled with `manage.py reextract_tables`."
    )


def format_row_line(i: int, row: dict) -> str:
    """One line for a missing-item list entry: value, page, and any name-like field."""
    value = row.get('_matched_value', 'N/A')
    page = row.get('_page', 'Unknown')
    line = f"{i}. `{value}` _(page {page})_"
    for key, val in row.items():
        if key.startswith('_') or not val or not str(val).strip():
            continue
        if any(kw in key.lower() for kw in ['nomenclature', 'designation', 'description', 'name']):
            line += f" — {val}"
            break
    return line


def format_comparison_response(comparison_result: dict) -> str:
    """Format comparison results as a single-pass summary: each number stated once,
    detail shown only for what needs a decision (missing / likely-matches)."""
    pdf1 = comparison_result.get('pdf1', 'File A')
    pdf2 = comparison_result.get('pdf2', 'File B')
    column = comparison_result.get('column', 'items')
    pdf1_total = comparison_result.get('pdf1_total', 0)
    pdf2_total = comparison_result.get('pdf2_total', 0)

    results_data = comparison_result.get('results', {})
    missing = results_data.get('in_pdf1_only', {})
    extra = results_data.get('in_pdf2_only', {})
    common = results_data.get('common', {})
    likely = results_data.get('likely_matches', {})
    confirmed = results_data.get('confirmed_matches', {})
    excluded = results_data.get('excluded', {})

    modified = results_data.get('modified', {})

    missing_count = missing.get('count', 0)
    extra_count = extra.get('count', 0)
    common_count = common.get('count', 0)
    likely_count = likely.get('count', 0)
    confirmed_count = confirmed.get('count', 0)
    excluded_count = excluded.get('count', 0)
    modified_count = modified.get('count', 0)

    # Column label: only show the "A ↔ B" pairing when the two files actually
    # use different header text — otherwise the identical-looking name (or a
    # smart-quote-only difference) just reads as visual noise.
    if ' ↔ ' in column:
        col1, col2 = column.split(' ↔ ', 1)
        col_line = (f"**Column checked:** `{col1}` in {pdf1}  ↔  `{col2}` in {pdf2}"
                    if col1.strip() != col2.strip() else f"**Column checked:** `{col1}`")
    else:
        col1 = column
        col_line = f"**Column checked:** `{column}`"

    response = f"## Comparison: {pdf1} vs {pdf2}\n\n"
    response += col_line + "\n\n"

    # ── One summary table, each count appears exactly once ──────────────────
    response += "| | Count |\n|---|---:|\n"
    response += f"| ✅ Matched exactly | {common_count} |\n"
    if confirmed_count:
        response += f"| ✔️ Confirmed by you (previously reviewed) | {confirmed_count} |\n"
    if likely_count:
        response += f"| ⚠️ Same item, written differently (review) | {likely_count} |\n"
    if modified_count:
        response += f"| 🔧 Matched, but other details differ (review) | {modified_count} |\n"
    missing_icon = "✅" if missing_count == 0 else "❌"
    response += f"| {missing_icon} In {pdf1} but missing from {pdf2} | {missing_count} |\n"
    if extra_count:
        response += f"| ℹ️ In {pdf2} but not in {pdf1} | {extra_count} |\n"
    response += "\n"

    # ── One bottom-line sentence — the actionable fact, stated once ─────────
    accounted = common_count + likely_count + confirmed_count
    if missing_count == 0:
        response += (
            f"**Bottom line:** all {pdf1_total} distinct `{col1}` values in {pdf1} "
            f"are accounted for in {pdf2}.\n\n"
        )
    else:
        response += (
            f"**Bottom line:** {pdf1} has {pdf1_total} distinct values in this column. "
            f"{accounted} are accounted for in {pdf2}"
            + (f" ({common_count} exact"
               + (f", {confirmed_count} confirmed" if confirmed_count else "")
               + (f", {likely_count} near-match" if likely_count else "") + ")")
            + f". **{missing_count} could not be found and need review.**\n\n"
        )
    response += "_Counts are of distinct values — a value repeated across pages counts once."
    if excluded_count:
        response += (
            f" {excluded_count} value(s) were generic placeholder text (e.g. "
            f"'Standard Item'), not real part identifiers, and are excluded from every "
            f"count above._\n\n---\n\n"
        )
    else:
        response += "_\n\n---\n\n"

    # ── Action needed: genuinely missing ─────────────────────────────────────
    if missing_count > 0:
        response += f"### ❌ Missing from {pdf2} ({missing_count})\n\n"
        rows = missing.get('rows', [])
        if rows:
            for i, row in enumerate(rows[:10], 1):
                response += format_row_line(i, row) + "\n"
        else:
            for i, v in enumerate(missing.get('values', [])[:10], 1):
                response += f"{i}. `{v}`\n"
        if missing_count > 10:
            response += f"\n_...and {missing_count - 10} more — full list in the Excel report._\n"
        response += "\n---\n\n"

    # ── Review needed: same item, but the documents disagree about it ────────
    if modified_count > 0:
        response += (
            f"### 🔧 Same item, different details ({modified_count})\n\n"
            f"These matched on `{col1}` but disagree elsewhere — the usual sign of a "
            f"revision that changed a quantity or a description.\n\n"
        )
        for item in modified.get('items', [])[:10]:
            label = f"`{item['value']}`"
            if item.get('matched_value'):
                label += f" (matched `{item['matched_value']}`)"
            response += f"\n**{label}**\n"
            for d in item['differences'][:4]:
                same_header = d['header_a'].lower() == d['header_b'].lower()
                where = (f"{d['header_a']}" if same_header
                         else f"{d['header_a']} → {d['header_b']}")
                response += f"- {where}: `{d['value_a']}` → `{d['value_b']}`\n"
        if modified_count > 10:
            response += f"\n_...and {modified_count - 10} more — full list in the Excel report._\n"
        response += "\n---\n\n"

    # ── Review needed: same item, different formatting ───────────────────────
    if likely_count > 0:
        response += (
            f"### ⚠️ Review — likely the same item, written differently ({likely_count})\n\n"
            "These weren't counted as missing, but weren't counted as matched either — "
            "confirm they're the same part before treating them either way.\n\n"
        )
        response += f"| {pdf1} | {pdf2} | Why | Times in {pdf2} |\n|---|---|---|---:|\n"
        for p in likely['pairs'][:15]:
            response += f"| `{p['pdf1_value']}` | `{p['pdf2_value']}` | {p['reason']} | {p.get('count_in_pdf2', '?')} |\n"
        if likely_count > 15:
            response += f"\n_...and {likely_count - 15} more — full list in the Excel report._\n"
        response += "\n---\n\n"

    # ── For reference only: no action implied ────────────────────────────────
    if common_count or extra_count or confirmed_count or excluded_count:
        response += "### ℹ️ For reference (no action needed)\n\n"
        if common_count:
            response += f"- **{common_count}** values matched exactly in both files.\n"
        if confirmed_count:
            response += (
                f"- **{confirmed_count}** values were confirmed as matches in an earlier "
                f"comparison (see below) and treated as matched here too.\n"
            )
        if extra_count:
            response += (
                f"- **{extra_count}** values appear in {pdf2} but weren't in {pdf1} — "
                f"expected if {pdf2} is the broader document (e.g. a full spares catalogue "
                f"versus a maintenance schedule); only worth checking if that's not the case here.\n"
            )
        if excluded_count:
            response += (
                f"- **{excluded_count}** values were generic placeholder text, not unique "
                f"part identifiers, and weren't compared at all.\n"
            )
        response += "\n_Full value lists for both are in the Excel report — see below._\n"

    # ── How many times each matched value repeats in the other file ──────────
    # Matching above is presence/absence on distinct values; this answers a
    # different question — a part can be "present" once but listed several
    # times in the other file (e.g. used in more than one assembly).
    occurrences = common.get('occurrences_in_pdf2', {})
    repeats = sorted(((v, c) for v, c in occurrences.items() if c > 1), key=lambda x: -x[1])
    if repeats:
        response += (
            f"\n---\n\n### 🔁 Appear more than once in {pdf2} ({len(repeats)})\n\n"
            f"These {len(repeats)} matched values are listed multiple times in {pdf2} "
            f"(counted once above, since presence was the question there):\n\n"
        )
        response += f"| {col1} | Times in {pdf2} |\n|---|---:|\n"
        for v, c in repeats[:15]:
            response += f"| `{v}` | {c} |\n"
        if len(repeats) > 15:
            response += f"\n_...and {len(repeats) - 15} more — full counts in the Excel report._\n"

    # ── Confirmed matches: what you previously told it were the same item ────
    if confirmed_count > 0:
        response += f"\n---\n\n### ✔️ Previously confirmed matches ({confirmed_count})\n\n"
        response += f"| {pdf1} | {pdf2} | Times in {pdf2} |\n|---|---|---:|\n"
        for p in confirmed['pairs'][:15]:
            response += f"| `{p['pdf1_value']}` | `{p['pdf2_value']}` | {p.get('count_in_pdf2', '?')} |\n"
        if confirmed_count > 15:
            response += f"\n_...and {confirmed_count - 15} more — full list in the Excel report._\n"

    return response


_STATUS_GLYPH = {"exact": "✅", "confirmed": "✔️", "likely": "⚠️", "missing": "❌", "n/a": "—"}


def format_multi_comparison_response(m: Dict[str, Any]) -> str:
    """Chat rendering of compare_many(): the actionable list first (missing from
    every file), then the presence matrix for partially-present values."""
    src, col = m["source"], m["column"]
    ok = m["ok_targets"]
    n_ok = len(ok)

    out = f"## Comparison: {src} vs {len(m['targets'])} file{'s' if len(m['targets']) != 1 else ''}\n\n"
    out += f"**Column checked:** `{col}` in {src}\n\n"

    if m["errors"]:
        out += "> ⚠️ Could not compare against: " + "; ".join(
            f"**{k}** ({v})" for k, v in m["errors"].items()) + "\n\n"
    if not ok:
        return out + "None of the target files could be compared."

    out += "| | Count |\n|---|---:|\n"
    out += f"| Distinct `{col}` values in {src} | {m['source_total']} |\n"
    out += f"| ✅ Accounted for in **all** {n_ok} file{'s' if n_ok != 1 else ''} | {len(m['present_everywhere'])} |\n"
    out += f"| ◐ Present in some, missing from others | {len(m['partial'])} |\n"
    out += f"| ❌ Missing from **every** file | {len(m['missing_everywhere'])} |\n\n"

    out += "**Per file:**\n\n| File | Column matched | Exact | Confirmed | Near-match | Missing |\n|---|---|---:|---:|---:|---:|\n"
    for name in ok:
        c = m["per_target"][name]
        out += f"| {name} | `{c['resolved_column']}` | {c['exact']} | {c['confirmed']} | {c['likely']} | {c['missing']} |\n"
    out += "\n_Counts are of distinct values. Near-matches are deterministic (leading zeros, annotation prefixes, typos in words) and listed for review — never silently counted as matched._\n\n---\n\n"

    if m["missing_everywhere"]:
        out += f"### ❌ Missing from every file ({len(m['missing_everywhere'])})\n\n"
        for i, v in enumerate(m["missing_everywhere"][:25], 1):
            out += f"{i}. `{v}`\n"
        if len(m["missing_everywhere"]) > 25:
            out += f"\n_...and {len(m['missing_everywhere']) - 25} more — full list in the Excel report._\n"
        out += "\n---\n\n"

    if m["partial"]:
        out += f"### ◐ Present in some files only ({len(m['partial'])})\n\n"
        out += "| Value | " + " | ".join(ok) + " |\n|---|" + "---|" * n_ok + "\n"
        for v in m["partial"][:25]:
            cells = []
            for name in ok:
                st = m["matrix"][v].get(name, "n/a")
                note = m["notes"].get(v, {}).get(name)
                cells.append(_STATUS_GLYPH.get(st, st) + (f" {note}" if note else ""))
            out += f"| `{v}` | " + " | ".join(cells) + " |\n"
        if len(m["partial"]) > 25:
            out += f"\n_...and {len(m['partial']) - 25} more — full matrix in the Excel report._\n"
        out += "\n"

    if m["present_everywhere"]:
        out += (f"\n_ℹ️ {len(m['present_everywhere'])} values are accounted for in every file "
                f"(no action needed) — full list in the Excel report._\n")
    return out


# ── Excel ──────────────────────────────────────────────────────────────────────

def generate_multi_comparison_excel(m: Dict[str, Any]) -> bytes:
    """Presence matrix + per-bucket sheets for compare_many()."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    fills = {
        "exact": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "confirmed": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "likely": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "missing": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    }

    def _header(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(horizontal="center")

    ws = wb.active
    ws.title = "Summary"
    _header(ws, ["Metric", "Value"])
    ws.append(["Source file", m["source"]])
    ws.append(["Column", m["column"]])
    ws.append(["Target files", ", ".join(m["targets"])])
    ws.append(["Distinct values in source", m["source_total"]])
    ws.append(["Accounted for in all files", len(m["present_everywhere"])])
    ws.append(["Present in some files only", len(m["partial"])])
    ws.append(["Missing from every file", len(m["missing_everywhere"])])
    for name, err in m["errors"].items():
        ws.append([f"Not compared: {name}", err])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 80

    ws = wb.create_sheet("Presence Matrix")
    ok = m["ok_targets"]
    _header(ws, [m["column"]] + ok + ["Status"])
    order = m["missing_everywhere"] + m["partial"] + m["present_everywhere"]
    for v in order:
        row = m["matrix"][v]
        status = ("missing everywhere" if v in m["missing_everywhere"]
                  else "present everywhere" if v in m["present_everywhere"] else "partial")
        cells = []
        for name in ok:
            st = row.get(name, "n/a")
            note = m["notes"].get(v, {}).get(name)
            cells.append(st + (f"  {note}" if note else ""))
        ws.append([v] + cells + [status])
        r = ws.max_row
        for j, name in enumerate(ok, start=2):
            f = fills.get(row.get(name))
            if f:
                ws.cell(row=r, column=j).fill = f
    ws.column_dimensions["A"].width = 28
    for j in range(2, len(ok) + 3):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 30
    ws.freeze_panes = "B2"

    ws = wb.create_sheet("Missing Everywhere")
    _header(ws, [m["column"]])
    for v in m["missing_everywhere"]:
        ws.append([v])
    ws.column_dimensions["A"].width = 32

    ws = wb.create_sheet("Per File")
    _header(ws, ["File", "Column matched", "Exact", "Confirmed", "Near-match", "Missing"])
    for name in ok:
        c = m["per_target"][name]
        ws.append([name, c["resolved_column"], c["exact"], c["confirmed"], c["likely"], c["missing"]])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 32

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def attach_excel_report(result: Dict[str, Any], report_type: str = "Comparison_Report",
                        excel_bytes: Optional[bytes] = None) -> Optional[str]:
    """Build the Excel for a comparison result, register it in the in-memory report
    store, and append the download link to result['formatted_answer'].
    Returns the job id (or None if the report could not be built)."""
    # The report store and the two-file Excel builder live in views.py; imported
    # lazily so this module stays importable from anywhere without a cycle.
    from . import views as _views
    try:
        if excel_bytes is None:
            excel_bytes = _views.generate_comparison_excel(result)
        job_id = _views.store_comparison_report(report_type, result, excel_bytes)
    except Exception as e:                       # a report failure must not kill the answer
        print(f"[COMPARISON] Excel report failed: {e}")
        return None
    result["formatted_answer"] = (result.get("formatted_answer", "") +
                                  f"\n\n---\n\n📥 **[Download Excel Report](/api/reports/download/{job_id}/)**")
    result["report_job_id"] = job_id
    return job_id


# ── Convenience: everything a chat router needs in one call ────────────────────

def _run_text_mode(thread, docs: List[Document], column: str,
                   prefer_literal: bool = False) -> Optional[Dict[str, Any]]:
    """Text mode as a complete chat answer, or None if the source column itself
    could not be resolved (in which case the caller's own error is the better one)."""
    m = compare_against_text(docs[0], docs[1:], column, thread=thread,
                             prefer_literal=prefer_literal)
    if not m.get("found"):
        return None
    m["formatted_answer"] = format_text_mode_response(m)
    try:
        attach_excel_report(m, report_type="Mention_Report",
                            excel_bytes=generate_text_mode_excel(m))
    except Exception as e:
        print(f"[COMPARISON] mention Excel failed: {e}")
    return m


def run_comparison(thread, docs: List[Document], column: str,
                   column_b: Optional[str] = None, comparison_type: str = "all",
                   match_any_column_in_b: bool = False, prefer_literal: bool = False,
                   requested_label: Optional[str] = None,
                   text_mode: bool = False) -> Dict[str, Any]:
    """
    Compare `docs` (2 → pairwise, ≥3 → source vs the rest) and return
    {found, formatted_answer, report_job_id?, result} ready for a chat reply.
    `found` is True even for the explanatory failure messages, so callers do not
    fall back to a narrative RAG answer for a question that was clearly a comparison.

    `text_mode=True` searches the targets' full text instead of a matching column
    (§ compare_against_text) — the only way to check a spares list against a manual.
    It is also used automatically when the column-to-column comparison fails purely
    because a target has no such column.
    """
    if len(docs) < 2:
        return {"found": False}

    if text_mode:
        answer = _run_text_mode(thread, docs, column, prefer_literal)
        if answer is not None:
            return answer
        # Source column unresolvable → fall through to the normal path, whose
        # "Column Not Found" message names the source's real headers.

    if len(docs) == 2:
        a, b = docs
        result = compare_documents(a, b, column, column_b=column_b,
                                   comparison_type=comparison_type,
                                   match_any_column_in_b=match_any_column_in_b,
                                   prefer_literal=prefer_literal, thread=thread)
        if result.get("found"):
            result["formatted_answer"] = format_comparison_response(result)
            attach_excel_report(result)
            return result

        # The pairwise compare failed. If the *source* column is fine and only the
        # target lacks one — the spares-list-vs-manual case — answer the question
        # that was actually asked by searching the target's text instead.
        if not text_mode:
            answer = _run_text_mode(thread, docs, column, prefer_literal)
            if answer is not None:
                answer["formatted_answer"] = (
                    f"> ℹ️ **{b.filename}** has no comparable `{column}` column, so I searched "
                    f"its full text for each value instead.\n\n" + answer["formatted_answer"])
                return answer

        requested = requested_label or (column if not column_b else
                                        f"{column} (in {a.filename}) / {column_b} (in {b.filename})")
        if result.get("error") == "no_data":
            msg = format_no_data(result, a.filename, b.filename)
        else:
            msg = format_column_not_found(
                result, requested, a.filename, b.filename,
                suggestions={a.filename: suggest_identifier_columns(a, thread),
                             b.filename: suggest_identifier_columns(b, thread)})
        return {"found": True, "formatted_answer": msg, "result": result}

    source, targets = docs[0], docs[1:]
    m = compare_many(source, targets, column, prefer_literal=prefer_literal, thread=thread)
    if not m["ok_targets"] and not text_mode:
        # No target had a comparable column (all manuals?) — search their text.
        answer = _run_text_mode(thread, docs, column, prefer_literal)
        if answer is not None:
            answer["formatted_answer"] = (
                f"> ℹ️ None of the target documents have a comparable `{column}` column, "
                f"so I searched their full text for each value instead.\n\n"
                + answer["formatted_answer"])
            return answer
    m["formatted_answer"] = format_multi_comparison_response(m)
    if m["found"]:
        try:
            attach_excel_report(m, report_type="Multi_Comparison_Report",
                                excel_bytes=generate_multi_comparison_excel(m))
        except Exception as e:
            print(f"[COMPARISON] multi Excel failed: {e}")
    m["found"] = True
    return m
